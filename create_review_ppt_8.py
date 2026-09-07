from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

OUT = 'PQC_Secure_Communication_Review_8_Slides_REVISED.pptx'
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BG=RGBColor(9,18,34); PANEL=RGBColor(18,32,55); PANEL2=RGBColor(24,43,70); WHITE=RGBColor(235,242,249); MUTED=RGBColor(166,187,207); CYAN=RGBColor(45,211,191); BLUE=RGBColor(72,151,255); GOLD=RGBColor(247,190,76); RED=RGBColor(244,106,106); GREEN=RGBColor(104,220,145)

def box(s,x,y,w,h,fill=PANEL,line=None,rounded=True):
    a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); a.fill.solid(); a.fill.fore_color.rgb=fill; a.line.color.rgb=line or fill
    if rounded: a.adjustments[0]=0.08
    return a

def txt(s,v,x,y,w,h,size=16,color=WHITE,bold=False,align=PP_ALIGN.LEFT,font='Aptos',valign=MSO_ANCHOR.TOP):
    a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); f=a.text_frame; f.clear(); f.word_wrap=True; f.margin_left=Inches(.06); f.margin_right=Inches(.06); f.vertical_anchor=valign; p=f.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=v; r.font.name=font; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color; return a

def bullets(s,items,x,y,w,h,size=15,color=WHITE): return txt(s,'\n'.join('- '+i for i in items),x,y,w,h,size,color)
def slide(title,section):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; box(s,0,0,13.333,.12,CYAN,rounded=False); txt(s,section.upper(),.62,.33,5.5,.25,10,CYAN,True); txt(s,title,.58,.66,11.9,.62,27,WHITE,True,font='Aptos Display'); txt(s,str(len(prs.slides)).zfill(2),12.35,.4,.45,.25,11,MUTED,True,PP_ALIGN.RIGHT); return s
def card(s,title,body,x,y,w,h,accent=CYAN,size=14): box(s,x,y,w,h,PANEL,PANEL2); box(s,x,y,.08,h,accent,rounded=False); txt(s,title,x+.2,y+.14,w-.3,.34,15,accent,True); txt(s,body,x+.2,y+.58,w-.3,h-.7,size,WHITE)
def flow(s,labels,y,x=.55,width=2.0,gap=.22):
    for i,label in enumerate(labels):
        xx=x+i*(width+gap); box(s,xx,y,width,.92,PANEL2,CYAN); txt(s,label,xx+.1,y+.13,width-.2,.62,13,WHITE,True,PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
        if i<len(labels)-1:
            l=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(xx+width),Inches(y+.46),Inches(xx+width+gap),Inches(y+.46)); l.line.color.rgb=CYAN; l.line.width=Pt(2)

# 1 Cover
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=BG; box(s,0,0,13.333,.16,CYAN,rounded=False); txt(s,'PQC SECURE COMMUNICATION',.72,1.18,10.5,.5,31,CYAN,True,font='Aptos Display'); txt(s,'Post-Quantum Cryptography Platform',.75,1.88,11,.65,35,WHITE,True,font='Aptos Display'); txt(s,'Project Review Presentation',.78,2.78,5.8,.35,19,GOLD,True); txt(s,'Secure Chat  |  Encrypted Mail  |  File Vault  |  Attack Lab',.78,3.4,10.9,.4,16,MUTED); box(s,.78,4.42,11.65,1.05,PANEL,PANEL2); txt(s,'Core objective',1.05,4.67,1.8,.25,13,CYAN,True); txt(s,'Demonstrate how post-quantum key establishment and authenticated encryption protect practical application workflows.',2.55,4.57,9.25,.5,17,WHITE); txt(s,'Stack: Flask | Socket.IO | SQLAlchemy | cryptography | liboqs',.8,6.48,10.5,.3,13,MUTED)

# 2
s=slide('Normal Application vs Our Application','Problem and solution'); card(s,'Normal secure application','Usually depends on RSA/ECC for key exchange and signatures.\n\nStrong against classical attackers, but a future quantum computer threatens these asymmetric assumptions.',.7,1.55,5.65,3.55,RED,16); card(s,'Our PQC application','Adds NIST-standardized post-quantum algorithms.\n\nHybrid mode combines X25519 + ML-KEM, then AES-256-GCM encrypts application data.',6.95,1.55,5.65,3.55,CYAN,16); txt(s,'Threat addressed',.8,5.55,2,.3,15,GOLD,True); txt(s,'Harvest Now, Decrypt Later: capture encrypted traffic today and decrypt it when quantum capability improves.',2.7,5.48,9.7,.45,16,WHITE)

# 3
s=slide('Application Structure and Section Creation','Architecture'); card(s,'Frontend','index.html\napp.js + quantum_flow.js\nChat, mail, files, labs, audit and key screens',.65,1.5,2.85,2.2,BLUE,14); card(s,'Flask app factory','create_app() initializes DB and Socket.IO, then registers blueprints for each section.',3.75,1.5,2.85,2.2,CYAN,14); card(s,'Application sections','/auth/*\n/api/*\n/api/mail/*\n/api/files/*\n/api/keys/*\n/api/audit/*',6.85,1.5,2.85,2.2,GOLD,14); card(s,'Core services','Crypto engine\nSQLAlchemy models\nSQLite persistence\nAudit + security monitor',9.95,1.5,2.7,2.2,GREEN,14); flow(s,['User action','Route / event','Crypto service','Database / room','Response'],4.55,width=2.25,gap=.2); txt(s,'A section is created as a Flask blueprint or Socket.IO event group, then registered from app/__init__.py.',.85,6.12,11.5,.35,14,MUTED,False,PP_ALIGN.CENTER)

# 4
s=slide('Main Algorithms and Their Roles','Cryptographic design'); rows=[('ML-KEM','Post-quantum key establishment','Lattice-based shared secret; Levels 512, 768, 1024.'),('X25519','Classical key exchange','Fast ECDH component in hybrid mode.'),('ML-DSA','Digital signatures','Authenticates handshake data and message ciphertext.'),('SLH-DSA','Hash-based signature option','Alternative stateless hash-based PQC signature.'),('AES-256-GCM','Actual message encryption','Encrypts chat, mail and files; tag detects tampering.'),('HKDF-SHA-384*','Final session-key derivation','Combines secrets and derives a 32-byte context-separated key.'),('Ascon-128a','Lightweight environments','Authenticated encryption for constrained/IoT use.'),('RSA-2048','Classical baseline','Comparison mode and attack demonstrations.')]
for i,(a,b,c) in enumerate(rows):
    y=1.34+i*.66; box(s,.65,y,2.05,.5,PANEL2,PANEL2); box(s,2.85,y,3.2,.5,PANEL,PANEL); box(s,6.25,y,6.4,.5,PANEL,PANEL); txt(s,a,.8,y+.12,1.75,.24,13,CYAN,True); txt(s,b,3,y+.12,2.9,.24,12,WHITE,True); txt(s,c,6.4,y+.1,5.95,.28,12,MUTED)
txt(s,'*Source-code correction: this project currently implements HKDF-SHA-384, not HKDF-SHA-256.',.75,6.75,11.8,.25,12,GOLD,True,PP_ALIGN.CENTER)

# 5
s=slide('How Session Creation and Handshake Work','Core flow'); flow(s,['Alice initiates','Bob accepts','Ephemeral keys','Sign + verify','Shared session key'],1.48,width=2.28,gap=.28); bullets(s,['User selects Classical, Modern Classical, PQC or Hybrid mode and a NIST level.','Hybrid: X25519 shared secret + ML-KEM shared secret.','Responder signs exchange material using ML-DSA; initiator verifies it.','HKDF-SHA-384 derives one 32-byte AES session key from the combined secret.','The session key stays in memory; a SHA-256 hash identifies the session in the database.'],.85,3.05,11.5,2.25,16); box(s,.95,5.78,11.3,.65,PANEL2,PANEL2); txt(s,'K_session = HKDF-SHA-384(X25519_secret || ML-KEM_secret)',1.2,5.96,10.8,.28,18,CYAN,True,PP_ALIGN.CENTER,font='Consolas')

# 6
s=slide('How Chat Messages Are Protected','Runtime workflow'); flow(s,['Plaintext','AES-256-GCM','Ciphertext + IV + tag','ML-DSA signature','Verify + deliver'],1.48,width=2.28,gap=.28); card(s,'Sender side','Encrypt with active 32-byte session key. Sign ciphertext and send sequence number, IV, tag and mode.',.75,3.05,3.7,2.05,BLUE,15); card(s,'Receiver side','Verify signature, check replay sequence, validate GCM tag/AAD, then deliver plaintext.',4.82,3.05,3.7,2.05,CYAN,15); card(s,'Security result','Wrong key or modified bytes fail authentication. Duplicate sequence numbers are rejected and logged.',8.89,3.05,3.7,2.05,GOLD,15); txt(s,'Stored Message fields: encrypted_payload, iv, auth_tag, signature, mode and sequence_number.',.8,6.05,11.5,.3,14,MUTED,False,PP_ALIGN.CENTER)

# 7
s=slide('Security Controls: Threat to Defense','Security model'); txt(s,'Every important security property is handled by a specific control in the application.',.75,1.38,11.6,.35,16,WHITE)
controls=[('Confidentiality','AES-256-GCM','Plaintext is converted to ciphertext; wrong keys cannot decrypt.'),('Integrity','GCM tag + SHA3','Bit changes and corrupted files are detected.'),('Authentication','ML-DSA / RSA','Only a valid identity signature is accepted.'),('Key safety','ML-KEM + X25519 + HKDF','Independent secrets become one session key.'),('Freshness','Sequence numbers','Previously accepted message sequences are rejected.'),('Accountability','AuditLog','Security actions record algorithm, result, risk and time.')]
for i,(a,b,c) in enumerate(controls):
    x=.7+(i%2)*6.15; y=1.95+(i//2)*1.25
    box(s,x,y,5.55,.92,PANEL,PANEL2); txt(s,a,x+.2,y+.14,1.45,.25,14,GOLD,True); txt(s,b,x+1.75,y+.14,1.8,.25,14,CYAN,True); txt(s,c,x+.2,y+.49,5.05,.24,12,MUTED)
box(s,.8,5.9,11.7,.55,PANEL2,PANEL2); txt(s,'Result: an attacker must bypass key exchange, identity verification, authenticated encryption and replay checks together.',1.05,6.06,11.2,.25,14,WHITE,True,PP_ALIGN.CENTER)

# 8
s=slide('Complete Working Flow and Review Takeaways','End-to-end summary'); flow(s,['Register + keys','Login','Handshake','Encrypt data','Verify + store','Audit / monitor'],1.4,x=.35,width=2.0,gap=.18); txt(s,'What happens in one secure chat session?',.8,2.85,4.4,.35,16,CYAN,True); bullets(s,['1. Registration creates a bcrypt credential hash and the user key suite.','2. Alice and Bob negotiate a mode; Hybrid combines X25519 and ML-KEM secrets.','3. ML-DSA authenticates the exchange; HKDF derives the 32-byte session key.','4. AES-256-GCM encrypts each message and binds metadata through AAD.','5. Receiver verifies signature, replay sequence and GCM tag before delivery.','6. Ciphertext and security events remain available for history and audit.'],.95,3.28,11.2,2.25,15)
box(s,.8,6.05,11.7,.55,PANEL2,PANEL2); txt(s,'Review conclusion: PQC protects the key-establishment stage; AEAD protects the actual application data.',1.05,6.2,11.2,.25,15,WHITE,True,PP_ALIGN.CENTER)

prs.save(OUT); print(OUT)
