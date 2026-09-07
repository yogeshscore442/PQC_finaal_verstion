from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_THEME_COLOR

OUT = 'PQC_Secure_Communication_Review.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(9, 18, 34)
PANEL = RGBColor(18, 32, 55)
PANEL2 = RGBColor(24, 43, 70)
WHITE = RGBColor(235, 242, 249)
MUTED = RGBColor(166, 187, 207)
CYAN = RGBColor(45, 211, 191)
BLUE = RGBColor(72, 151, 255)
GOLD = RGBColor(247, 190, 76)
RED = RGBColor(244, 106, 106)
GREEN = RGBColor(104, 220, 145)


def rect(slide, x, y, w, h, fill=PANEL, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False, align=PP_ALIGN.LEFT, font='Aptos', valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06); tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return box


def bullets(slide, items, x, y, w, h, size=16, color=WHITE, gap=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True; tf.margin_left = Inches(0.08); tf.margin_right = Inches(0.04)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = item; p.level = 0; p.space_after = Pt(gap); p.font.name = 'Aptos'; p.font.size = Pt(size); p.font.color.rgb = color
        p.text = '- ' + item
    return box


def base(title, kicker=None):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    rect(s, 0, 0, 13.333, 0.12, CYAN)
    if kicker: text(s, kicker.upper(), 0.62, 0.34, 5.4, 0.28, 10, CYAN, True)
    text(s, title, 0.58, 0.67, 12.1, 0.62, 27, WHITE, True, font='Aptos Display')
    text(s, str(len(prs.slides)).zfill(2), 12.35, 0.42, 0.45, 0.3, 11, MUTED, True, PP_ALIGN.RIGHT)
    return s


def card(slide, title, body, x, y, w, h, accent=CYAN, size=14):
    rect(slide, x, y, w, h, PANEL, PANEL2, True)
    rect(slide, x, y, 0.08, h, accent)
    text(slide, title, x+0.22, y+0.16, w-0.35, 0.35, 15, accent, True)
    text(slide, body, x+0.22, y+0.62, w-0.35, h-0.78, size, WHITE)


def flow(slide, labels, y=3.1, x=0.7, boxw=2.15, gap=0.35):
    for i, label in enumerate(labels):
        xx = x + i*(boxw+gap)
        rect(slide, xx, y, boxw, 1.0, PANEL2, CYAN, True)
        text(slide, label, xx+0.12, y+0.18, boxw-0.24, 0.62, 14, WHITE, True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        if i < len(labels)-1:
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xx+boxw), Inches(y+0.5), Inches(xx+boxw+gap), Inches(y+0.5))
            line.line.color.rgb = CYAN; line.line.width = Pt(2)

# 1
s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
rect(s, 0, 0, 13.333, 0.16, CYAN)
text(s, 'PQC SECURE COMMUNICATION', 0.72, 1.15, 8.5, 0.62, 31, CYAN, True, font='Aptos Display')
text(s, 'Post-Quantum Cryptography Platform', 0.75, 1.93, 9.7, 0.7, 36, WHITE, True, font='Aptos Display')
text(s, 'Project Review Presentation', 0.78, 2.82, 5.8, 0.4, 19, GOLD, True)
text(s, 'Secure Chat  |  Quantum-Resistant Mail  |  Encrypted File Vault  |  Attack Simulation', 0.78, 3.42, 10.9, 0.48, 15, MUTED)
rect(s, 0.78, 4.42, 11.65, 1.12, PANEL, PANEL2, True)
text(s, 'Core idea', 1.05, 4.67, 1.5, 0.3, 13, CYAN, True)
text(s, 'Use NIST-standardized PQC primitives with authenticated encryption to protect communication from current and future cryptanalytic threats.', 2.35, 4.58, 9.55, 0.56, 17, WHITE)
text(s, 'Implementation: Flask + Socket.IO + SQLAlchemy + cryptography + liboqs', 0.8, 6.55, 9, 0.3, 13, MUTED)

# 2
s=base('Review Map', 'Presentation roadmap')
flow(s, ['Problem', 'Architecture', 'Crypto', 'Features', 'Security', 'Testing'], y=2.05, x=0.7, boxw=1.75, gap=0.35)
card(s, 'Review goal', 'Explain what the platform solves, how data moves through it, which algorithms are used, and how the defenses are verified.', 0.8, 4.1, 5.75, 1.55, CYAN, 16)
card(s, 'Demo storyline', 'Register -> login -> start hybrid handshake -> send encrypted chat -> test tampering/replay -> inspect audit logs.', 6.8, 4.1, 5.75, 1.55, GOLD, 16)

# 3
s=base('The Problem: Classical Crypto Is Not Future-Proof', 'Motivation')
card(s, 'Today', 'RSA and ECC protect many systems using integer factorization and discrete logarithm hardness.', 0.7, 1.75, 3.85, 1.65, BLUE, 15)
card(s, 'Quantum threat', 'Shor\'s algorithm can theoretically break these asymmetric assumptions on a sufficiently capable quantum computer.', 4.75, 1.75, 3.85, 1.65, RED, 15)
card(s, 'HNDL attack', 'Harvest Now, Decrypt Later: capture encrypted data today and decrypt it when quantum capability arrives.', 8.8, 1.75, 3.85, 1.65, GOLD, 15)
text(s, 'Project response', 0.75, 4.15, 2.2, 0.35, 15, CYAN, True)
bullets(s, ['Adopt NIST-standardized post-quantum key establishment and signatures.', 'Use hybrid key establishment during migration: X25519 + ML-KEM.', 'Use AES-256-GCM for fast authenticated bulk encryption.', 'Demonstrate attacks and defenses with real cryptographic operations.'], 0.85, 4.58, 11.5, 1.65, 17)

# 4
s=base('What the Platform Provides', 'Objectives and scope')
items=[('Real-time chat','1-to-1 and group messaging through Socket.IO.'),('Secure mail','Encrypted body plus sender signature and inbox/sent/read flows.'),('File vault','Per-file AES-GCM key, encrypted storage, integrity digest.'),('Attack Lab','Controlled wrong-key, tamper, signature, replay and MITM demonstrations.'),('Crypto Lab','Primitive matrix, real execution checks and host-dependent benchmarks.'),('Auditability','Security events, operation results, risk levels and timestamps.')]
for i,(a,b) in enumerate(items):
    x=0.72+(i%3)*4.15; y=1.55+(i//3)*2.25
    card(s,a,b,x,y,3.65,1.65,[CYAN,BLUE,GOLD,GREEN,RED,CYAN][i],15)

# 5
s=base('System Architecture', 'End-to-end view')
card(s,'Frontend','HTML/CSS/JavaScript dashboard\nSocket.IO client\nChart.js + Three.js visualizations',0.7,1.65,3.0,2.05,BLUE,15)
card(s,'Application server','Flask app factory\nBlueprint routes\nSocket.IO event handlers\nSession-based authentication',4.05,1.65,3.0,2.05,CYAN,15)
card(s,'Crypto engine','cryptography primitives\nliboqs PQC wrappers\nHKDF key derivation\nAEAD encryption',7.4,1.65,2.35,2.05,GOLD,15)
card(s,'Persistence','SQLite / SQLAlchemy\nCiphertext records\nAudit logs\nSession metadata',10.1,1.65,2.35,2.05,GREEN,15)
flow(s,['User action','HTTP / REST','Socket.IO','Crypto + DB','Verified response'],y=4.65,x=0.9,boxw=2.15,gap=0.3)
text(s,'App factory registers /auth, /api, /api/mail, /api/files, /api/keys and /api/audit blueprints.',0.9,6.12,11.5,0.35,14,MUTED)

# 6
s=base('Identity Registration and Authentication', 'Step 1: establish identity')
flow(s,['Register user','Hash password','Generate key suite','Store identity','Login session'],y=1.72,x=0.65,boxw=2.25,gap=0.3)
text(s,'During registration',0.75,3.22,2.5,0.35,16,CYAN,True)
bullets(s,['Username is normalized and duplicate usernames are rejected.', 'Password is stored as a bcrypt hash, never as plaintext.', 'RSA-2048, ML-DSA, SLH-DSA and X25519 key pairs are generated.', 'Public/private key material is stored in the User record for this prototype.', 'Registration and login success/failure are written to AuditLog.'],0.9,3.68,11.3,1.75,16)
rect(s,0.8,5.78,11.6,0.55,PANEL2,PANEL2,True)
text(s,'Review point: authentication identity and cryptographic identity are created together.',1.05,5.92,11.1,0.25,15,GOLD,True,PP_ALIGN.CENTER)

# 7
s=base('Cryptographic Toolbox', 'Algorithms and roles')
rows=[('Key establishment','ML-KEM-512 / 768 / 1024','Post-quantum KEM; NIST FIPS 203'),('Modern classical','X25519','Fast ECDH baseline'),('Digital signatures','ML-DSA-44 / 65 / 87','PQC identity authentication; FIPS 204'),('Backup signature','SLH-DSA','Hash-based signature; FIPS 205'),('Bulk encryption','AES-256-GCM / ChaCha20-Poly1305','Confidentiality + integrity'),('Lightweight AEAD','Ascon-128a','Constrained/IoT-oriented option'),('Hash + KDF','SHA3 family, SHAKE-256, HKDF-SHA-384','Integrity, context binding, key separation')]
for i,(a,b,c) in enumerate(rows):
    y=1.48+i*0.68
    rect(s,0.7,y,2.4,0.53,PANEL2,PANEL2,True); rect(s,3.2,y,3.5,0.53,PANEL,PANEL,True); rect(s,6.85,y,5.75,0.53,PANEL,PANEL,True)
    text(s,a,0.85,y+0.13,2.1,0.25,13,CYAN,True); text(s,b,3.35,y+0.13,3.2,0.25,13,WHITE,True); text(s,c,7.0,y+0.13,5.35,0.25,13,MUTED)

# 8
s=base('Hybrid Key Establishment', 'Core security concept')
text(s,'Why hybrid?',0.75,1.5,2.0,0.35,17,GOLD,True)
text(s,'Two independent mechanisms contribute to one session key. The design supports migration: classical performance plus PQC protection.',0.75,1.92,11.5,0.48,16,WHITE)
flow(s,['X25519 ECDH','ML-KEM encapsulation','Concatenate secrets','HKDF-SHA-384','32-byte session key'],y=3.0,x=0.65,boxw=2.25,gap=0.3)
rect(s,0.9,5.0,11.4,0.85,PANEL,PANEL2,True)
text(s,'K_session = HKDF-SHA-384(X25519_secret || ML-KEM_secret, info = hybrid-v1)',1.15,5.28,10.9,0.3,18,CYAN,True,PP_ALIGN.CENTER,font='Consolas')
text(s,'The same derived key is stored in memory for both communication directions; only a SHA-256 session hash is persisted.',0.9,6.18,11.4,0.3,13,MUTED,False,PP_ALIGN.CENTER)

# 9
s=base('Interactive Chat Handshake', 'Step 2: agree on a session key')
flow(s,['Alice initiates','Bob accepts','Generate ephemeral keys','Sign + verify','Activate session'],y=1.58,x=0.65,boxw=2.25,gap=0.3)
bullets(s,['Modes: Classical, Modern Classical, PQC and Hybrid.', 'PQC mode uses ML-KEM for shared secret and ML-DSA for authentication.', 'Hybrid mode signs responder X25519 public key plus KEM ciphertext.', 'Signature verification protects handshake material from substitution/MITM.', 'Both users receive the same session hash after successful derivation.'],0.9,3.35,11.5,1.8,16)
rect(s,0.9,5.72,11.4,0.55,PANEL2,PANEL2,True); text(s,'Failure path: invalid identity signature -> handshake_failed -> no active session.',1.1,5.86,11.0,0.25,15,RED,True,PP_ALIGN.CENTER)

# 10
s=base('Message Protection Pipeline', 'Step 3: encrypt, authenticate, deliver')
flow(s,['Plaintext','AES-256-GCM','Ciphertext + IV + tag','ML-DSA signature','Verify + decrypt'],y=1.62,x=0.65,boxw=2.25,gap=0.3)
text(s,'What is checked before delivery?',0.8,3.3,4.0,0.35,16,CYAN,True)
bullets(s,['The sender is authenticated through the active session and user identity.', 'The message signature is verified against ciphertext bytes.', 'Sequence number is checked against the session replay set.', 'GCM authentication tag validates ciphertext and associated data.', 'Only after checks pass is plaintext emitted to the recipient and ciphertext stored.'],0.95,3.75,11.2,1.7,16)
text(s,'AAD binds context such as mode and sequence number, so metadata changes also cause authentication failure.',0.9,6.05,11.4,0.35,14,GOLD,False,PP_ALIGN.CENTER)

# 11
s=base('Secure Mail Flow', 'Asynchronous encrypted communication')
flow(s,['Compose mail','Create Email ID','HKDF context key','AES-GCM encrypt','ML-DSA sign','Read: verify + decrypt'],y=1.55,x=0.45,boxw=1.95,gap=0.22)
text(s,'Context-bound derivation',0.75,3.2,3.0,0.35,16,CYAN,True)
text(s,'Seed = mail-ID-sender-receiver-subject',0.75,3.65,5.2,0.35,18,WHITE,True,font='Consolas')
bullets(s,['Mail body is stored as encrypted_body with IV, authentication tag and signature.', 'Default non-Classical signing path uses ML-DSA; Classical mode uses RSA.', 'Read route checks recipient/sender access, verifies signature and performs GCM decryption.', 'Inbox, sent, read state and audit events are supported.'],6.35,3.45,6.1,1.6,15)
rect(s,0.8,5.75,11.6,0.55,PANEL2,PANEL2,True); text(s,'Implementation note: this is server-side encrypted mail; the server performs read-time decryption.',1.0,5.88,11.2,0.26,14,GOLD,True,PP_ALIGN.CENTER)

# 12
s=base('Encrypted File Vault', 'Confidentiality + integrity')
flow(s,['Upload file','Random file key','AES-256-GCM','Wrap key','Store .enc','Download + verify'],y=1.5,x=0.45,boxw=1.95,gap=0.22)
bullets(s,['Allowed document, image, text and archive extensions; maximum size is 25 MB.', 'Each upload receives a fresh 32-byte per-file AES key.', 'The file key is wrapped by a platform vault master key derived through HKDF.', 'Stored payload is ciphertext under instance/uploads/*.enc.', 'Integrity choices: SHA3-256 default, SHA3-512 or SHAKE-256 optional.', 'Download unwraps key, decrypts, checks digest and returns the file.'],0.85,3.2,11.5,2.1,15)
text(s,'Design boundary: download creates plaintext in temp_downloads for response, so the vault is encrypted storage, not strict zero-knowledge storage.',0.85,5.85,11.5,0.48,14,GOLD,False,PP_ALIGN.CENTER)

# 13
s=base('Attack Lab: Demonstrable Defenses', 'Security verification')
att=[('Wrong key','AES-GCM raises InvalidTag','Confidentiality'),('Bit tamper','Tag validation fails','Integrity'),('Modified payload','ML-DSA verify = False','Authenticity'),('Replay','Duplicate sequence rejected','Freshness'),('MITM substitution','Signature check rejects altered keys','Handshake identity'),('Plain HTTP baseline','Shows why encryption is needed','Awareness')]
for i,(a,b,c) in enumerate(att):
    x=0.7+(i%3)*4.15; y=1.55+(i//3)*2.0
    card(s,a,b+'\n\nDefense: '+c,x,y,3.65,1.45,[RED,GOLD,CYAN,BLUE,GREEN,MUTED][i],14)
text(s,'The strongest demonstrations invoke real crypto operations and observe the rejection or verification result.',0.8,6.05,11.5,0.35,14,MUTED,False,PP_ALIGN.CENTER)

# 14
s=base('Data Model and Persistence', 'What the database remembers')
entities=['User\ncredentials + key suite','Message\nciphertext + tag + signature','Email\nencrypted body + signature','Attachment\nwrapped file key + digest','GroupChat / Member\ngroup access','AuditLog\nsecurity events','BenchmarkResult\ntiming + size','UserSessionKey\nhash + status']
for i,e in enumerate(entities):
    x=0.65+(i%4)*3.15; y=1.65+(i//4)*2.15
    rect(s,x,y,2.75,1.25,PANEL,PANEL2,True); text(s,e,x+0.12,y+0.22,2.5,0.8,15,WHITE,True,PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
text(s,'Security principle in schema: Message and Email records persist encrypted payload fields, not readable plaintext bodies.',0.8,6.05,11.5,0.35,14,CYAN,False,PP_ALIGN.CENTER)

# 15
s=base('API and Frontend Experience', 'How users interact')
card(s,'Authentication','/auth/register\n/auth/login\n/auth/logout\n/auth/me',0.7,1.55,2.75,2.0,BLUE,14)
card(s,'Crypto + lab','/api/crypto/primitives\n/api/crypto/test_primitive\nbenchmarks and attack endpoints',3.7,1.55,2.75,2.0,GOLD,14)
card(s,'Communication','Socket.IO chat events\n/api/mail/*\n/api/files/*\nGroup chat APIs',6.7,1.55,2.75,2.0,CYAN,14)
card(s,'Operations','/api/keys/*\n/api/audit/*\nsecurity monitor\nonline status',9.7,1.55,2.75,2.0,GREEN,14)
text(s,'Dashboard views',0.8,4.25,2.2,0.35,16,CYAN,True)
bullets(s,['Authentication and identity generation', 'Chat, groups and online status', 'Secure mail and encrypted files', 'Attack Lab, Crypto Lab and benchmarks', 'Audit logs and key/session management'],0.95,4.7,11.1,1.45,16)

# 16
s=base('Testing and Verification Strategy', 'Evidence')
card(s,'Unit crypto tests','AES-GCM, ChaCha20, Ascon, SHA3, SHAKE, HKDF, RSA, X25519, ML-KEM, ML-DSA, SLH-DSA and hybrid equality.',0.7,1.55,3.75,2.0,CYAN,14)
card(s,'Integration tests','Registration/login, mail send/read, file upload/download, primitive matrix, benchmarks and attack endpoints.',4.78,1.55,3.75,2.0,BLUE,14)
card(s,'Interactive tests','Two Socket.IO clients complete a Hybrid handshake, compare session hashes and exchange an encrypted message.',8.85,1.55,3.75,2.0,GOLD,14)
text(s,'Important assertions',0.8,4.22,3.0,0.35,16,CYAN,True)
bullets(s,['Encapsulated and decapsulated secrets are equal.', 'Modified ciphertext is rejected by AEAD.', 'Modified signed data is rejected by ML-DSA.', 'Wrong key produces authentication failure.', 'NIST security levels 1, 3 and 5 are exercised.'],0.95,4.68,11.2,1.55,16)

# 17
s=base('Suggested Live Review Demo', 'Five-minute walkthrough')
flow(s,['Register Alice/Bob','Login','Hybrid handshake','Send message','Attack Lab','Audit log'],y=1.55,x=0.55,boxw=1.95,gap=0.22)
for i,(a,b) in enumerate([('1','Create two users and show generated key fingerprints.'),('2','Select Hybrid + NIST Level 3; accept handshake and show same session hash.'),('3','Send a message; show encrypted packet fields and recipient plaintext.'),('4','Run bit tamper / wrong key / replay tests; show blocked results.'),('5','Open audit monitor and explain algorithm, result and risk level.')]):
    y=3.25+i*0.52; text(s,a,0.95,y,0.35,0.3,14,CYAN,True); text(s,b,1.45,y,10.7,0.3,14,WHITE)

# 18
s=base('Current Limitations and Honest Scope', 'Engineering review')
items=[('Server-side decryption','Mail and file download decrypt on the server; this is not strict end-to-end or zero-knowledge design.'),('Key storage','Private identity keys are stored in database fields without an additional at-rest encryption layer.'),('File ownership','Authenticated list/download paths should be tightened with explicit attachment ownership and authorization checks.'),('Replay model','Live chat rejects duplicate sequence numbers, but does not fully enforce monotonic order or timestamps.'),('Prototype deployment','Use production secrets, HTTPS/WSS, secure cookies, rate limits, secret management and a production DB before deployment.')]
for i,(a,b) in enumerate(items):
    y=1.42+i*0.87; rect(s,0.75,y,11.85,0.68,PANEL,PANEL2,True); text(s,a,1.0,y+0.18,2.75,0.25,14,GOLD,True); text(s,b,3.85,y+0.15,8.35,0.34,13,WHITE)

# 19
s=base('Future Roadmap', 'From academic prototype to hardened product')
flow(s,['Protect private keys','True client E2E','Authorization hardening','Deploy securely','Scale + monitor'],y=1.62,x=0.65,boxw=2.25,gap=0.3)
bullets(s,['Encrypt private keys with a user-controlled KMS/HSM-backed wrapping key.', 'Move mail/file key derivation and decryption to trusted clients where appropriate.', 'Enforce attachment ownership, group authorization and least privilege on every route.', 'Use HTTPS/WSS, secure cookie flags, CSRF protection, rate limiting and environment secrets.', 'Replace in-memory session/replay state with a shared, durable store for multi-worker deployment.', 'Add migration/versioning, observability, threat modeling and external security review.'],0.9,3.35,11.4,2.25,15)

# 20
s=base('Conclusion', 'Key takeaways')
text(s,'A practical PQC learning and demonstration platform',0.8,1.55,11.7,0.55,25,CYAN,True,PP_ALIGN.CENTER,font='Aptos Display')
card(s,'1. Understand','It shows why current asymmetric crypto needs a quantum migration path.',0.8,2.65,3.7,1.5,BLUE,15)
card(s,'2. Implement','It combines ML-KEM, ML-DSA, X25519, HKDF and AES-GCM in usable communication workflows.',4.82,2.65,3.7,1.5,CYAN,15)
card(s,'3. Verify','It includes real crypto tests and controlled attack simulations instead of only theoretical claims.',8.84,2.65,3.7,1.5,GOLD,15)
rect(s,1.0,5.05,11.25,0.85,PANEL2,PANEL2,True)
text(s,'Review closing line: “The platform connects cryptographic theory to visible, testable application behavior.”',1.3,5.31,10.65,0.35,18,WHITE,True,PP_ALIGN.CENTER)
text(s,'Thank you',0.8,6.45,11.7,0.3,16,MUTED,False,PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
